"""PySide6 main window for configuring and executing the access analysis."""

from __future__ import annotations
import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING

import cartopy.crs as ccrs
from PySide6.QtCore import QObject, QThread, QTimer, Signal
from PySide6.QtGui import QIcon
from PySide6.QtWidgets import (
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from src.models import (
    AnalysisConfig,
    AnalysisOptions,
    AnalysisResult,
    GroundStationConfig,
    GroundTrackPoint,
    OrbitConfig,
    PassStatistic,
    PropagationConfig,
    ScenarioConfig,
)
from src.services.access_analysis import run_access_analysis
from src.ui.constants import HIST_BIN_OPTIONS
from src.ui.plot_helpers import PlotHelpersMixin
from src.ui.tabs import (
    GroundTabMixin,
    LinkBudgetTabMixin,
    MissionTabMixin,
    OrbitSummaryTabMixin,
    VisualizationTabMixin,
)

if TYPE_CHECKING:
    import numpy as np
    from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
    from src.ui.opengl import GlobeWidget
    from pyqtgraph import PlotWidget
    from PySide6.QtGui import QColor
    from PySide6.QtWidgets import (
        QComboBox,
        QListWidget,
        QProgressBar,
        QSlider,
    )



class AnalysisCancelledException(Exception):
    """Raised internally to signal that the analysis was cancelled by the user."""


class AnalysisWorker(QObject):
    """Background worker that executes the analysis and emits progress."""

    progress = Signal(float)
    finished = Signal(object)
    error = Signal(str)

    def __init__(
        self,
        config: AnalysisConfig,
        stations: list[GroundStationConfig],
    ) -> None:
        super().__init__()
        self._config = config
        self._stations = stations
        self._cancelled = False

    def run(self) -> None:
        try:
            result = run_access_analysis(
                self._config,
                self._stations,
                progress_callback=self._emit_progress,
            )
        except AnalysisCancelledException:
            # Graceful cancellation: no result to emit, just stop the worker.
            self.error.emit("Analysis cancelled by user.")
            return
        except Exception as exc:  # pragma: no cover - GUI execution path
            self.error.emit(str(exc))
            return
        self.finished.emit(result)

    def _emit_progress(self, value: float) -> None:
        # Allow cancellation to be picked up from within the Orekit sampling loop.
        if getattr(self, "_cancelled", False):
            raise AnalysisCancelledException()
        self.progress.emit(value)


class GroundStationApp(
    GroundTabMixin,
    MissionTabMixin,
    VisualizationTabMixin,
    LinkBudgetTabMixin,
    OrbitSummaryTabMixin,
    PlotHelpersMixin,
    QMainWindow,
):
    """Main PyQt window that exposes the configuration controls."""

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Cosmic Comms Toolbox")
        icon_path = (
            Path(__file__).resolve().parents[2] / "resources" / "img" / "menu_icon.png"
        )
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))
        self._station_presets: list[GroundStationConfig] = []
        self._is_dirty = True
        self._current_config: AnalysisConfig | None = None
        self._last_result: AnalysisResult | None = None
        self._histogram_bin_seconds = HIST_BIN_OPTIONS[1]
        self._histogram_bar_data: list[tuple[float, float, float]] = []
        self._station_color_map: dict[str, QColor] = {}
        self.station_table: QTableWidget | None = None
        self.select_all_button: QPushButton | None = None
        self.clear_selection_button: QPushButton | None = None
        self.map_canvas: FigureCanvasQTAgg | None = None
        self.map_axes = None
        self.map_projection = ccrs.PlateCarree()
        self.map_status_label: QLabel | None = None
        self._map_station_artists: list = []
        self.analysis_context_label: QLabel | None = None
        self.pie_canvas: FigureCanvasQTAgg | None = None
        self.pie_axes = None
        self.visual_globe_widget: GlobeWidget | None = None
        self.mission_globe_widget: GlobeWidget | None = None
        self.mission_globe_status_label: QLabel | None = None
        self._analysis_thread: QThread | None = None
        self._analysis_worker: AnalysisWorker | None = None
        self._pending_config: AnalysisConfig | None = None
        self.run_progress: QProgressBar | None = None
        self.stop_button: QPushButton | None = None
        self.link_budget_station_combo: QComboBox | None = None
        self.link_budget_table: QTableWidget | None = None
        self.link_budget_summary_label: QLabel | None = None
        self._link_budget_auto_enabled = False
        self._latest_access_series: dict | None = None
        self._link_budget_rate_curve: tuple[np.ndarray, np.ndarray] | None = None
        self._latest_station_rate_series: dict[str, np.ndarray] | None = None
        self._latest_combined_rate_series: np.ndarray | None = None
        self._downlink_total_label: QLabel | None = None
        self._downlink_per_orbit_label: QLabel | None = None
        self.pass_volume_plot = None
        self.daily_volume_plot = None
        self.pass_bin_label: QLabel | None = None
        self.daily_bin_label: QLabel | None = None
        self._pass_volume_bin_width: float | None = None
        self._daily_volume_bin_width: float | None = None
        self._last_pass_volume_samples = None
        self._last_daily_volume_samples = None
        self._active_station_lookup: dict[str, GroundStationConfig] = {}
        self.visual_station_tabs: QTabWidget | None = None
        self.visual_pass_status_label: QLabel | None = None
        self.visual_play_button: QPushButton | None = None
        self.visual_time_slider: QSlider | None = None
        self.visual_speed_slider: QSlider | None = None
        self.visual_view_tabs: QTabWidget | None = None
        self.visual_graph_combo: QComboBox | None = None
        self.visual_graph_plot: PlotWidget | None = None
        self._visual_graph_data: dict[str, np.ndarray] | None = None
        self._visual_station_lists: dict[str, QListWidget] = {}
        self._visual_selected_pass: PassStatistic | None = None
        self._visual_pass_track: list[GroundTrackPoint] | None = None
        self._visual_animation_index: int = 0
        self._visual_station_ecef: tuple[float, float, float] | None = None
        self._visual_contact_window: tuple[datetime, datetime] | None = None
        self._visual_animation_timer = QTimer(self)
        self._visual_animation_timer.setInterval(33)
        self._visual_animation_timer.timeout.connect(
            self._advance_visualization_animation
        )
        self._visual_animation_timer.setSingleShot(False)
        self._visual_base_step = 0.25
        self._visual_speed_multiplier = 1.0
        self._visual_anim_fraction = 0.0
        self._visual_frame_mode: str = "ECI"
        self._visual_reference_epoch: datetime | None = None
        self._visual_freq_listener_connected = False
        self.mission_frame_combo: QComboBox | None = None
        self.mission_window_slider: QSlider | None = None
        self.mission_window_label: QLabel | None = None
        self._mission_frame_mode: str = "ECI"
        self._mission_reference_epoch: datetime | None = None
        self._mission_orbit_period_s: float = 0.0
        self._mission_window_fractions: tuple[float, float] = (0.0, 1.0)
        self._mission_ground_track: list[GroundTrackPoint] | None = None
        self._mission_globe_refresh_timer = QTimer(self)
        self._mission_globe_refresh_timer.setSingleShot(True)
        self._mission_globe_refresh_timer.setInterval(120)
        self._mission_globe_refresh_timer.timeout.connect(self._refresh_mission_globe)
        self._build_ui()

    def _set_run_button_state(self, state: str) -> None:
        """Update run button text and color based on state."""
        if not hasattr(self, "run_button"):
            return
        palette = {
            "dirty": ("Run Analysis", "#FFA500", "#FFD580"),
            "running": ("Running…", "#1E90FF", "#87CEFA"),
            "success": ("Completed", "#2E8B57", "#90EE90"),
        }
        text, color, hover = palette.get(state, palette["dirty"])
        self.run_button.setText(text)
        self.run_button.setStyleSheet(
            f"QPushButton {{ background-color: {color}; border-radius: 4px; padding: 6px; }}"
            f"QPushButton:hover {{ background-color: {hover}; }}"
        )

    def _mark_dirty(self) -> None:
        """Mark configuration as needing re-run."""
        self._is_dirty = True
        self._set_run_button_state("dirty")

    def _build_contact_analysis_tab(self) -> QTabWidget:
        """Create the Contact Analysis tab with Statistics, Visualization, and Link Budget sub-tabs."""
        contact_tabs = QTabWidget()
        contact_tabs.setTabPosition(QTabWidget.TabPosition.North)

        # Build the individual components
        gs_stats_tab = self._build_analysis_tabs()
        visualization_tab = self._build_visualization_tab()
        link_budget_tab = self._build_link_budget_tab()
        data_volume_tab = self._build_data_volume_tab()
        self._attach_visualization_frequency_listener()

        # Add them as sub-tabs
        contact_tabs.addTab(gs_stats_tab, "Contact Statistics")
        contact_tabs.addTab(visualization_tab, "Pass Visualization")
        contact_tabs.addTab(link_budget_tab, "Link Budget Tool")
        contact_tabs.addTab(data_volume_tab, "Data Volume")

        return contact_tabs

    def _build_ui(self) -> None:
        """Compose the widgets and layouts making up the window."""
        central_widget = QWidget(self)
        self.setCentralWidget(central_widget)
        root_layout = QVBoxLayout(central_widget)
        self.main_tabs = QTabWidget()
        root_layout.addWidget(self.main_tabs)
        ground_station_tab = self._build_ground_station_tab()
        mission_tab = self._build_mission_analysis_tab()
        contact_analysis_tab = self._build_contact_analysis_tab()
        orbit_summary_tab = self._build_orbit_summary_tab()
        self.main_tabs.addTab(ground_station_tab, "Ground Station Selection")
        self.main_tabs.addTab(mission_tab, "Mission Configuration")
        self.main_tabs.addTab(contact_analysis_tab, "Contact Analysis")
        self.main_tabs.addTab(orbit_summary_tab, "Orbit Summary")



    def _build_analysis_tabs(self) -> QTabWidget:
        """Create the Statistics/Contact Analysis tab stack."""
        self.summary_label = QLabel("No analysis run yet.")
        self.analysis_context_label = QLabel("Active stations: manual entry")
        self.analysis_context_label.setWordWrap(True)
        self.results_table = QTableWidget(0, 6)
        self.results_table.setHorizontalHeaderLabels(
            [
                "Pass",
                "Station",
                "AOS (UTC)",
                "LOS (UTC)",
                "Duration (min)",
                "Max Elev (deg)",
            ]
        )
        self.results_table.horizontalHeader().setStretchLastSection(True)

        # Create export buttons
        export_layout = QHBoxLayout()
        export_csv_button = QPushButton("Export to CSV")
        export_csv_button.clicked.connect(lambda: self._export_results_table("csv"))
        export_xlsx_button = QPushButton("Export to XLSX")
        export_xlsx_button.clicked.connect(lambda: self._export_results_table("xlsx"))
        export_layout.addWidget(export_csv_button)
        export_layout.addWidget(export_xlsx_button)
        export_layout.addStretch()

        stats_tab = QWidget()
        stats_layout = QVBoxLayout(stats_tab)
        stats_layout.addWidget(self.summary_label)
        stats_layout.addWidget(self.analysis_context_label)
        stats_layout.addLayout(export_layout)
        stats_layout.addWidget(self.results_table, stretch=1)
        tabs = QTabWidget()
        tabs.setTabPosition(QTabWidget.TabPosition.North)
        tabs.addTab(stats_tab, "Statistics")
        for title, widget in self._create_contact_plot_tabs():
            tabs.addTab(widget, title)
        return tabs




    def _handle_run_clicked(self) -> None:
        """Build the analysis configuration and execute it."""
        self._set_run_button_state("running")
        perform_ground_passes = getattr(self, "ground_pass_checkbox", None)
        ground_pass_enabled = (
            perform_ground_passes.isChecked()
            if perform_ground_passes is not None
            else True
        )

        stations: list[GroundStationConfig] = []
        primary_station: GroundStationConfig | None = None

        if ground_pass_enabled:
            stations = self._collect_active_stations_for_run()
            self._active_station_lookup = {station.name: station for station in stations}
            if not stations:
                self._set_run_button_state("dirty")
                return
            primary_station = stations[0]
        else:
            # No ground-station analysis → no stations required.
            self._active_station_lookup = {}
            primary_station = None
        try:
            config = self._build_config_from_inputs(
                primary_station=primary_station,
                enable_ground_pass_analysis=ground_pass_enabled,
            )
        except ValueError as exc:  # pragma: no cover - GUI validation
            QMessageBox.warning(self, "Invalid Input", str(exc))
            self._set_run_button_state("dirty")
            return
        self._clear_visualization_display("Running analysis…")
        self._show_run_progress_ui()
        self._start_analysis_worker(config, stations)

    def _show_run_progress_ui(self) -> None:
        """Display the in-button progress bar state."""
        if self.run_progress:
            self.run_progress.setValue(0)
            self.run_progress.setFormat("0%")
            self.run_progress.show()
        self.run_button.setEnabled(False)
        self.run_button.hide()
        if getattr(self, "stop_button", None) is not None:
            self.stop_button.setEnabled(True)
            self.stop_button.show()

    def _hide_run_progress_ui(self) -> None:
        """Restore the run button once analysis completes."""
        if self.run_progress:
            self.run_progress.hide()
        self.run_button.show()
        self.run_button.setEnabled(True)
        if getattr(self, "stop_button", None) is not None:
            self.stop_button.setEnabled(False)

    def _start_analysis_worker(
        self, config: AnalysisConfig, stations: list[GroundStationConfig]
    ) -> None:
        """Kick off the background analysis worker."""
        self._pending_config = config
        self._analysis_thread = QThread(self)
        self._analysis_worker = AnalysisWorker(config, stations)
        self._analysis_worker.moveToThread(self._analysis_thread)
        self._analysis_thread.started.connect(self._analysis_worker.run)
        self._analysis_worker.progress.connect(self._update_run_progress)
        self._analysis_worker.finished.connect(self._handle_analysis_success)
        self._analysis_worker.error.connect(self._handle_analysis_error)
        self._analysis_worker.finished.connect(self._analysis_thread.quit)
        self._analysis_worker.error.connect(self._analysis_thread.quit)
        self._analysis_worker.finished.connect(self._analysis_worker.deleteLater)
        self._analysis_worker.error.connect(self._analysis_worker.deleteLater)
        self._analysis_thread.finished.connect(self._cleanup_analysis_thread)
        self._analysis_thread.start()

    def _handle_stop_clicked(self) -> None:
        """Request cancellation of the currently running analysis, if any."""
        worker = getattr(self, "_analysis_worker", None)
        if worker is not None:
            setattr(worker, "_cancelled", True)

    def _cleanup_analysis_thread(self) -> None:
        """Release worker references after the thread stops."""
        self._analysis_thread = None
        self._analysis_worker = None

    def _update_run_progress(self, value: float) -> None:
        """Update the progress indicator."""
        if not self.run_progress:
            return
        clamped = max(0, min(100, int(value)))
        self.run_progress.setValue(clamped)
        self.run_progress.setFormat(f"{clamped:.0f}%")

    def _handle_analysis_success(self, result: AnalysisResult) -> None:
        """Handle successful completion of the analysis."""
        self._current_config = self._pending_config
        self._pending_config = None
        if self._current_config is not None:
            self._visual_reference_epoch = self._current_config.scenario.start_time
        self._populate_results_table(result)
        self._update_summary_label(result)
        self._last_result = result
        self._set_run_button_state("success")
        self._update_plots(result)
        self._mission_reference_epoch = self._visual_reference_epoch
        self._mission_orbit_period_s = float(result.orbit_period_seconds or 0.0)
        self._mission_ground_track = result.ground_track
        self._mission_window_fractions = (0.0, 1.0)
        if self.mission_window_slider:
            self.mission_window_slider.setValue(
                self.mission_window_slider.maximum()
            )
        self._update_mission_window_label()
        self._refresh_mission_globe()
        self._refresh_visualization_pass_tabs(result)
        self._store_access_series(result)
        self._update_downlink_summary()
        self._update_orbit_summary(result)
        # If SSO is enabled, keep the inclination field synced after the run.
        if hasattr(self, "_apply_sso_if_enabled"):
            try:
                self._apply_sso_if_enabled()  # type: ignore[attr-defined]
            except Exception:
                pass
        self._update_run_progress(100.0)
        self._hide_run_progress_ui()

    def _handle_analysis_error(self, message: str) -> None:
        """Handle errors raised by the analysis worker."""
        self._pending_config = None
        self._set_run_button_state("dirty")
        self._hide_run_progress_ui()
        # Treat user-initiated cancellation as a non-fatal condition.
        if message.strip().lower().startswith("analysis cancelled"):
            return
        QMessageBox.critical(self, "Analysis Error", message)

    def _build_config_from_inputs(
        self,
        primary_station: GroundStationConfig | None,
        *,
        enable_ground_pass_analysis: bool,
    ) -> AnalysisConfig:
        """Translate widget state into a strongly-typed config."""
        start_dt = self._qdatetime_to_utc(self.start_datetime)
        end_dt = self._qdatetime_to_utc(self.end_datetime)
        if end_dt <= start_dt:
            raise ValueError("End time must be later than the start time.")
        if enable_ground_pass_analysis and primary_station is None:
            raise ValueError(
                "Select at least one ground station when ground-station analysis is enabled."
            )
        ground = primary_station
        altitude_km = float(self.altitude_input.value()) if getattr(self, "altitude_input", None) is not None else 0.0
        ecc = float(self.ecc_input.value())
        # Derive SMA from geodetic altitude above WGS-84.
        sma_km = float(getattr(self, "_WGS84_EQUATORIAL_RADIUS_KM", 6378.137)) + altitude_km

        # If SSO is enabled, override inclination using orbital mechanics.
        inc_deg = float(self.inc_input.value())
        if getattr(self, "sso_checkbox", None) is not None and self.sso_checkbox.isChecked():
            computed = self._compute_sso_inclination_deg(altitude_km=altitude_km, eccentricity=ecc)  # type: ignore[attr-defined]
            if computed is not None:
                inc_deg = float(computed)
                # Keep the UI in sync with the applied inclination.
                try:
                    self.inc_input.blockSignals(True)
                    self.inc_input.setValue(inc_deg)
                finally:
                    self.inc_input.blockSignals(False)
        orbit = OrbitConfig(
            semi_major_axis_km=sma_km,
            eccentricity=ecc,
            inclination_deg=inc_deg,
            raan_deg=self.raan_input.value(),
            arg_perigee_deg=self.argp_input.value(),
            mean_anomaly_deg=self.mean_anom_input.value(),
        )
        propagation = PropagationConfig(
            # Mean-IC solver is always applied, which requires numerical propagation.
            propagator_type="numerical",
            # Pass detection is intentionally ungated by any minimum elevation mask.
            min_elevation_deg=0.0,
            sample_step_seconds=float(self.sample_step_input.value()),
            enable_drag=self.drag_checkbox.isChecked(),
            use_mean_ic_finder=True,
            mean_ic_target_altitude_km=float(altitude_km),
            drag_area_m2=float(self.drag_area_input.value()),
            drag_cd=float(self.drag_cd_input.value()),
        )
        scenario = ScenarioConfig(start_time=start_dt, end_time=end_dt)
        options = AnalysisOptions(
            compute_ground_station_passes=enable_ground_pass_analysis
        )
        return AnalysisConfig(
            ground_station=ground,
            orbit=orbit,
            propagation=propagation,
            scenario=scenario,
            options=options,
        )

    def _qdatetime_to_utc(self, widget) -> datetime:
        """Convert a QDateTime widget value into a timezone-aware datetime."""
        qdt = widget.dateTime()
        py_dt = None
        for attr in ("toPyDateTime", "toPython"):
            converter = getattr(qdt, attr, None)
            if callable(converter):
                py_dt = converter()
                break
        if py_dt is None:
            # Fallback to epoch conversion if direct Python conversion is unavailable.
            py_dt = datetime.fromtimestamp(qdt.toSecsSinceEpoch(), tz=timezone.utc)
        if py_dt.tzinfo is None:
            py_dt = py_dt.replace(tzinfo=timezone.utc)
        else:
            py_dt = py_dt.astimezone(timezone.utc)
        return py_dt

    def _populate_results_table(self, result) -> None:
        """Render the per-pass statistics inside the table widget."""
        passes = sorted(result.passes, key=lambda item: item.aos)
        self.results_table.setRowCount(len(passes))
        for row, item in enumerate(passes):
            display_index = row + 1
            self.results_table.setItem(row, 0, QTableWidgetItem(str(display_index)))
            station_name = item.station_name or "Ground Station"
            self.results_table.setItem(row, 1, QTableWidgetItem(station_name))
            self.results_table.setItem(
                row, 2, QTableWidgetItem(item.aos.strftime("%d-%b-%Y %H:%M:%S"))
            )
            self.results_table.setItem(
                row, 3, QTableWidgetItem(item.los.strftime("%d-%b-%Y %H:%M:%S"))
            )
            self.results_table.setItem(
                row, 4, QTableWidgetItem(f"{item.duration_minutes:.2f}")
            )
            self.results_table.setItem(
                row, 5, QTableWidgetItem(f"{item.max_elevation_deg:.1f}")
            )
        self.results_table.resizeColumnsToContents()

    def _export_results_table(self, format_type: str) -> None:
        """Export the results table to CSV or XLSX format."""
        if self.results_table.rowCount() == 0:
            QMessageBox.warning(
                self, "No Data", "No results to export. Please run an analysis first."
            )
            return

        # Generate timestamp for filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        outputs_dir = Path(__file__).resolve().parents[2] / "outputs"
        outputs_dir.mkdir(exist_ok=True)

        # Collect headers
        headers = []
        for col in range(self.results_table.columnCount()):
            header_item = self.results_table.horizontalHeaderItem(col)
            headers.append(header_item.text() if header_item else f"Column {col}")

        # Collect data rows
        data = []
        for row in range(self.results_table.rowCount()):
            row_data = []
            for col in range(self.results_table.columnCount()):
                item = self.results_table.item(row, col)
                row_data.append(item.text() if item else "")
            data.append(row_data)

        try:
            if format_type == "csv":
                filename = outputs_dir / f"contact_statistics_{timestamp}.csv"
                with open(filename, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(headers)
                    writer.writerows(data)
                QMessageBox.information(
                    self, "Export Successful", f"Data exported to:\n{filename}"
                )
            elif format_type == "xlsx":
                try:
                    import openpyxl
                    from openpyxl.styles import Font

                    filename = outputs_dir / f"contact_statistics_{timestamp}.xlsx"
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Contact Statistics"

                    # Write headers with bold font
                    for col_idx, header in enumerate(headers, start=1):
                        cell = ws.cell(row=1, column=col_idx, value=header)
                        cell.font = Font(bold=True)

                    # Write data
                    for row_idx, row_data in enumerate(data, start=2):
                        for col_idx, value in enumerate(row_data, start=1):
                            ws.cell(row=row_idx, column=col_idx, value=value)

                    # Auto-adjust column widths
                    for col in ws.columns:
                        max_length = 0
                        column = col[0].column_letter
                        for cell in col:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        adjusted_width = min(max_length + 2, 50)
                        ws.column_dimensions[column].width = adjusted_width

                    wb.save(filename)
                    QMessageBox.information(
                        self, "Export Successful", f"Data exported to:\n{filename}"
                    )
                except ImportError:
                    QMessageBox.critical(
                        self,
                        "Missing Dependency",
                        "openpyxl is required for XLSX export.\n"
                        "Install it with: pip install openpyxl",
                    )
        except Exception as e:
            QMessageBox.critical(
                self, "Export Failed", f"Failed to export data:\n{str(e)}"
            )

    def _update_summary_label(self, result) -> None:
        """Show the aggregated statistics to the user."""
        # If ground-station analysis was disabled, show a dedicated message.
        if (
            getattr(self, "_current_config", None) is not None
            and getattr(self._current_config, "options", None) is not None
            and not self._current_config.options.compute_ground_station_passes
        ):
            self.summary_label.setText(
                "Ground-station pass analysis was disabled for this run."
            )
            if self.analysis_context_label is not None:
                self.analysis_context_label.setText("Ground-station analysis disabled.")
            return

        summary = result.summary
        base_text = (
            f"Passes: {summary.total_passes} | Total Access: {summary.total_access_minutes:.1f} min | "
            f"Coverage: {summary.coverage_percent:.2f}% | Avg: {summary.avg_duration_minutes:.2f} min | "
            f"Min: {summary.min_duration_minutes:.2f} min | Max: {summary.max_duration_minutes:.2f} min"
        )
        mean_ic_report = getattr(result, "mean_ic_report", None)
        if isinstance(mean_ic_report, str) and mean_ic_report.strip():
            base_text += f" | {mean_ic_report}"
        station_summaries = getattr(result, "station_summaries", [])
        if len(station_summaries) > 1:
            breakdown = ", ".join(
                f"{entry.station_name}: {entry.total_access_minutes:.1f} min"
                for entry in station_summaries
            )
            base_text += f" | Stations: {breakdown}"
        elif station_summaries:
            entry = station_summaries[0]
            base_text += f" | Station: {entry.station_name}"
        self.summary_label.setText(base_text)
        if self.analysis_context_label is not None:
            if station_summaries:
                names_list = [entry.station_name for entry in station_summaries]
                display_names = names_list[:5]
                if len(names_list) > 5:
                    display_names.append("…")
                names = ", ".join(display_names)
                self.analysis_context_label.setText(
                    f"Active stations ({len(station_summaries)}): {names}"
                )
            else:
                self.analysis_context_label.setText("Active stations: manual entry")
